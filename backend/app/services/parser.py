from pathlib import Path

import pymupdf
from pptx import Presentation

from app.services.chunker import ExtractedSection


SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".md", ".txt"}


class UnsupportedDocumentError(ValueError):
    pass


def parse_document(path: Path) -> list[ExtractedSection]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix in {".md", ".txt"}:
        return _parse_text(path)
    raise UnsupportedDocumentError(f"Unsupported file type: {suffix}")


def _parse_pdf(path: Path) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    with pymupdf.open(path) as document:
        for index, page in enumerate(document, start=1):
            sections.append(ExtractedSection(text=page.get_text("text"), page=index))
    return sections


def _parse_pptx(path: Path) -> list[ExtractedSection]:
    presentation = Presentation(path)
    sections: list[ExtractedSection] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        sections.append(ExtractedSection(text="\n".join(texts), slide=index))
    return sections


def _parse_text(path: Path) -> list[ExtractedSection]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [ExtractedSection(text=text, section=path.name)]
